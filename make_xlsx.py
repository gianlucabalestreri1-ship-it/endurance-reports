#!/usr/bin/env python3
"""Generador de Excel de Endurance Tool. Tres modos:

  python3 make_xlsx.py                    # diario: report/feedback-data.json -> report/Endurance-Feedback.xlsx
  python3 make_xlsx.py expiry             # semanal: report/plans-expiry-data.json -> report/Planes-Vencimientos.xlsx
  python3 make_xlsx.py dossier <data.json># dossier: report/dossier-<plan>-data.json -> report/Dossier-<Atleta>.xlsx

Sin fórmulas (export estático): no requiere recalc de LibreOffice.
Requiere openpyxl:  pip install -q openpyxl
"""
import json, sys, re, unicodedata
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

A = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="F15A24")
HEAD_FONT = Font(name=A, bold=True, color="FFFFFF", size=11)
CELL = Font(name=A, size=10); BOLD = Font(name=A, size=10, bold=True)
HOT = Font(name=A, size=11, bold=True, color="C00000")
thin = Side(style="thin", color="D9D9D9"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top"); CTR = Alignment(horizontal="center", vertical="top"); TOP = Alignment(vertical="top")
F_UNREAD = PatternFill("solid", fgColor="E7F7EE"); F_HOT = PatternFill("solid", fgColor="FDE7E7"); F_PAIN = PatternFill("solid", fgColor="FFF1DD")
F_IMPROVE = PatternFill("solid", fgColor="EBE0F7")   # lila: progresión (buena) en un dolor/molestia
TIPO_FILL = {"red": PatternFill("solid", fgColor="E5484D"), "yellow": PatternFill("solid", fgColor="D9A400"),
             "blue": PatternFill("solid", fgColor="3B82F6"), "gray": PatternFill("solid", fgColor="9AA0AA")}
TIPO_LABEL = {"HIT": "HIT", "ThT": "ThT", "LIT": "LIT", "FZA": "Fuerza", "TEST": "Test", "OTRO": "—"}


def header(ws, ncols, r=1, freeze=True):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=r, column=c); cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.border = BORDER; cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 26
    if freeze:
        ws.freeze_panes = "A" + str(r + 1)   # string, no materializa fila fantasma


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def tipo_cell(cell, intensity):
    color = (intensity or {}).get("color", "gray")
    hexc = ((intensity or {}).get("hex") or "").lstrip("#")
    cell.fill = PatternFill("solid", fgColor=hexc) if len(hexc) == 6 else TIPO_FILL.get(color, TIPO_FILL["gray"])
    cell.font = Font(name=A, size=9, bold=True, color="FFFFFF" if color != "yellow" else "333333")
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def drop_ghost_row(path, sheet):
    wb2 = load_workbook(path); ws2 = wb2[sheet]
    if all(ws2.cell(row=2, column=c).value in (None, "") for c in range(1, 15)):
        ws2.delete_rows(2, 1); ws2.freeze_panes = "A2"; wb2.save(path)


def slug(s):
    s = unicodedata.normalize("NFD", s or "")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-") or "atleta"


# ---------------------------------------------------------------------------
# Modo diario (default)
# ---------------------------------------------------------------------------

def build_daily():
    data = json.load(open("report/feedback-data.json"))
    rows = data["sessionRows"]; fit = data["fit"]; comps = data["competitions"]; cds = data["countdowns"]
    upcoming = data.get("upcoming") or []
    # Hoja Feedback: agrupar por atleta (alfabético, sin acentos/mayúsc.) y por fecha,
    # para no saltar de un atleta a otro y mantener contexto. La priorización por
    # dolor/sRPE vive en la hoja "⚠ Atención".
    def _akey(r):
        n = unicodedata.normalize("NFD", r.get("athlete", "") or "")
        n = "".join(c for c in n if unicodedata.category(c) != "Mn").lower().strip()
        return (n, r.get("date", "") or "", r.get("moment", "") or "")
    rows = sorted(rows, key=_akey)
    OUT = "report/Endurance-Feedback.xlsx"
    wb = Workbook()

    # ---- Feedback ----  (agrupado por atleta; cada tanda abre con una fila-separador
    #                       combinada a lo ancho con el nombre centrado, y cierra con una
    #                       fila vacía de respiro; sin columna "Atleta")
    ws = wb.active; ws.title = "Feedback"
    H = ["Fecha", "Mom.", "Sesión", "Deporte", "Tipo", "sRPE", "Feedback / comentario", "Dolor/molestia", "FIT (laps + métrica)"]
    NC = len(H)
    ws.append(H); header(ws, NC); widths(ws, [12, 7, 24, 11, 18, 7.43, 52, 12, 30])
    SEP_FILL = PatternFill("solid", fgColor="C9D6E5")
    SEP_FONT = Font(name=A, bold=True, size=12, color="1F3A5F")
    CMID = Alignment(horizontal="center", vertical="center", wrap_text=True)   # contenido: centrado + medio + ajustado
    cur = None
    for r in rows:
        ath = (r.get("athlete") or "—").strip()
        if ath != cur:
            if cur is not None:                     # fila vacía de respiro entre atletas
                ws.append([""] * NC)
            cur = ath
            ws.append([ath] + [""] * (NC - 1)); br = ws.max_row   # fila-separador con el nombre
            ws.merge_cells(start_row=br, start_column=1, end_row=br, end_column=NC)
            sc0 = ws.cell(row=br, column=1); sc0.value = ath; sc0.fill = SEP_FILL; sc0.font = SEP_FONT
            sc0.alignment = CMID; sc0.border = BORDER
        hot = (r.get("rpe") or 0) >= 6
        pc = r.get("painClass") or ("active" if r.get("pain") else "none")
        pain = (pc == "active"); improving = (pc == "improving")
        inten = r.get("intensity", {}) or {}; cat = inten.get("cat", "OTRO")
        tipo = TIPO_LABEL.get(cat, cat) + ((" · " + r["sessionType"]) if r.get("sessionType") else "")
        # FIT: resumen (km · dur · FC) + laps/vueltas
        fit_lines = []
        if r.get("hasFit"):
            summ = [r.get("fitDistance") or "", r.get("fitDuration") or ""]
            if r.get("fitHr"): summ.append("FC" + str(r.get("fitHr")))
            summ = " · ".join(p for p in summ if p)
            if summ: fit_lines.append(summ)
            for L in (r.get("fitLaps") or []):
                fit_lines.append("L%s %s %.2fkm %s %s/%s" % (L.get("n"), L.get("time", ""), L.get("distKm", 0),
                                                             L.get("pace", ""), L.get("hrA", ""), L.get("hrM", "")))
        fitv = "\n".join(fit_lines) if fit_lines else "—"
        dolor_val = "SÍ" if pain else ("Mejora 🟣" if improving else "NO")
        ws.append([r.get("date", ""), r.get("moment", ""), r.get("session", ""), r.get("sport", ""), tipo,
                   r.get("rpe"), (r.get("comment") or "").strip(), dolor_val, fitv])
        i = ws.max_row
        rf = F_PAIN if pain else (F_IMPROVE if improving else (F_HOT if hot else (F_UNREAD if r.get("unread") else None)))
        for c in range(1, NC + 1):
            cell = ws.cell(row=i, column=c); cell.font = CELL if c != 9 else Font(name=A, size=8)
            cell.border = BORDER; cell.alignment = CMID
            if rf and c != 5: cell.fill = rf
        tipo_cell(ws.cell(row=i, column=5), inten)                                # Tipo
        ws.cell(row=i, column=6).font = HOT if hot else BOLD                      # sRPE
        if pain: ws.cell(row=i, column=8).font = Font(name=A, size=10, bold=True, color="C00000")      # Dolor
        elif improving: ws.cell(row=i, column=8).font = Font(name=A, size=10, bold=True, color="6B3FA0")  # Mejora (lila)

    # ---- ⚠ Atención ----
    wa = wb.create_sheet("⚠ Atención")
    wa["A1"] = "REPORTE — PRESTAR ATENCIÓN"; wa["A1"].font = Font(name=A, bold=True, size=13, color="C00000")
    wa["A2"] = f"Coach {data.get('coach','')} · generado {data.get('generatedAt','')} · notificaciones sin leer"
    wa["A2"].font = Font(name=A, size=9, color="808080")
    n_hot = len([1 for r in rows if (r.get("rpe") or 0) >= 6]); n_pain = len([1 for r in rows if r.get("pain")]); n_fit = len([1 for r in rows if r.get("hasFit")])
    ct = data.get("counts", {})
    kpis = [
        ("Sin leer — total (todas las secciones)", ct.get("totalUnread", "")),
        ("Atletas DISTINTOS con feedback pendiente", ct.get("athletes", "")),
        ("Feedbacks de sesión sin leer", ct.get("atletasNotifs", len(rows))),
        ("  (filas al agrupar por sesión)", len(rows)),
        ("Competencias hechas", len(comps)),
        ("Próximas competencias", len(upcoming) or len(cds)),
        ("Con FIT", n_fit),
        ("sRPE >= 6", n_hot),
        ("Con dolor/molestia", n_pain),
    ]
    for r, (lbl, val) in enumerate(kpis, start=4):
        wa.cell(row=r, column=1, value=lbl).font = Font(name=A, size=10)
        vc = wa.cell(row=r, column=2, value=val); vc.alignment = CTR
        vc.font = HOT if lbl in ("sRPE >= 6", "Con dolor/molestia") else Font(name=A, size=11, bold=True)
    wa.cell(row=4 + len(kpis), column=1,
            value='Nota: "Atletas" es la sección/pestaña del panel (feedbacks de sesión), no la cantidad de atletas.').font = Font(name=A, size=8, italic=True, color="808080")
    start = 4 + len(kpis) + 3
    for i, h in enumerate(["Motivo", "Atleta", "Fecha", "Sesión / Competición", "Tipo", "sRPE", "Detalle"], 1):
        c = wa.cell(row=start, column=i, value=h); c.fill = HEAD_FILL; c.font = HEAD_FONT; c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)
    wa.freeze_panes = wa.cell(row=start + 1, column=1)
    widths(wa, [16, 20, 12, 26, 10, 7, 60])
    att = []
    def _pc(x): return x.get("painClass") or ("active" if x.get("pain") else "none")
    for r in rows:
        pc = _pc(r); reasons = []
        if (r.get("rpe") or 0) >= 6: reasons.append("sRPE≥6")
        if pc == "active": reasons.append("Dolor")
        elif pc == "improving": reasons.append("Molestia ↓ (mejora)")
        if reasons:
            att.append(("+".join(reasons), r.get("athlete", ""), r.get("date", ""), r.get("session", ""),
                        TIPO_LABEL.get((r.get("intensity") or {}).get("cat", "OTRO")), r.get("rpe"),
                        (r.get("comment") or "").strip()[:200], pc))
    for c in comps:
        att.append(("Competencia ✓", c.get("athlete", ""), c.get("date", ""), c.get("competition") or c.get("session", ""),
                    "Comp " + (c.get("priority") or ""), c.get("rpe"), (c.get("comment") or "").strip()[:200], _pc(c)))
    att = sorted(att, key=lambda x: (0 if "Dolor" in x[0] else (2 if "mejora" in x[0] else 1), -(x[5] or 0)))
    rr = start + 1
    for motivo, ath, fecha, sess, tipo, rpe, det, pc in att:
        for i, v in enumerate([motivo, ath, fecha, sess, tipo, rpe, det], 1):
            cell = wa.cell(row=rr, column=i, value=v); cell.font = CELL; cell.border = BORDER
            cell.alignment = WRAP if i == 7 else (CTR if i in (3, 5, 6) else TOP)
        fill = F_PAIN if pc == "active" else (F_IMPROVE if pc == "improving" else F_HOT)
        for i in range(1, 8):
            if i != 5: wa.cell(row=rr, column=i).fill = fill
        motivo_color = "6B3FA0" if pc == "improving" else "C00000"
        wa.cell(row=rr, column=1).font = Font(name=A, size=9, bold=True, color=motivo_color)
        wa.cell(row=rr, column=6).font = HOT if (rpe or 0) >= 6 else BOLD
        rr += 1

    # ---- Competencias ----
    wc = wb.create_sheet("Competencias")
    H3 = ["", "Atleta", "Fecha", "Competición", "Prio", "sRPE", "Feedback", "Estado"]
    wc.append(H3); header(wc, len(H3)); widths(wc, [4, 20, 12, 30, 6, 7, 50, 14])
    comps = sorted(comps, key=lambda c: (_akey({"athlete": c.get("athlete", "")})[0], c.get("date", "") or ""))
    for c in comps:
        wc.append(["●" if c.get("unread") else "", c.get("athlete", ""), c.get("date", ""),
                   c.get("competition") or c.get("session", ""), c.get("priority", ""), c.get("rpe"),
                   (c.get("comment") or "").strip(), "✓ Completada"])
        i = wc.max_row
        for k in range(1, len(H3) + 1):
            cell = wc.cell(row=i, column=k); cell.font = CELL; cell.border = BORDER
            cell.alignment = WRAP if k == 7 else (CTR if k in (1, 3, 5, 6) else TOP)
        wc.cell(row=i, column=6).font = HOT if (c.get("rpe") or 0) >= 6 else BOLD
        wc.cell(row=i, column=8).font = Font(name=A, size=10, bold=True, color="2E9E5B")

    # ---- Próximas competencias (agenda + cuenta regresiva) ----
    base = wc.max_row + 3
    wc.cell(row=base, column=2, value=f"PRÓXIMAS COMPETENCIAS ({len(upcoming) or len(cds)})").font = Font(name=A, bold=True, size=12, color="F15A24")
    hdr = ["Atleta", "Competición", "Fecha", "Aviso", "Prio"]
    for i, h in enumerate(hdr, 2):
        c = wc.cell(row=base + 1, column=i, value=h); c.fill = HEAD_FILL; c.font = HEAD_FONT
        c.border = BORDER; c.alignment = Alignment(vertical="center")
    rr = base + 2
    rows_up = upcoming if upcoming else [
        {"athlete": c.get("athlete", ""), "competition": c.get("competition", ""), "date": c.get("date", ""),
         "aviso": c.get("title", ""), "priority": c.get("priority", "A")} for c in cds]
    rows_up = sorted(rows_up, key=lambda u: (_akey({"athlete": u.get("athlete", "")})[0], u.get("date", "") or ""))
    for u in rows_up:
        aviso = u.get("aviso", "") or ("esta semana" if u.get("source") == "agenda" else "")
        if u.get("avisadoEl"):
            aviso = f"{aviso} · aviso {u['avisadoEl']}"
        vals = [u.get("athlete", ""), u.get("competition", "") or "—", u.get("date", "") or "—",
                aviso, u.get("priority", "")]
        for i, v in enumerate(vals, 2):
            cell = wc.cell(row=rr, column=i, value=v); cell.font = CELL; cell.border = BORDER; cell.alignment = TOP
        if u.get("source") == "agenda":
            for i in range(2, 7): wc.cell(row=rr, column=i).fill = F_UNREAD
            wc.cell(row=rr, column=4).font = BOLD
        pcell = wc.cell(row=rr, column=6); pcell.alignment = CTR
        if (u.get("priority") or "") == "A": pcell.font = Font(name=A, size=10, bold=True, color="C00000")
        rr += 1

    wb.save(OUT)   # sin drop_ghost_row: el freeze por string ya evita la fila fantasma, y delete_rows corrompe celdas combinadas
    print(f"saved {OUT} | feedback {len(rows)} | atención {len(att)} | competencias {len(comps)} | próximas {len(rows_up)}")


# ---------------------------------------------------------------------------
# Modo expiry (informe semanal de vencimientos)
# ---------------------------------------------------------------------------

def build_expiry():
    data = json.load(open("report/plans-expiry-data.json"))
    OUT = "report/Planes-Vencimientos.xlsx"
    wb = Workbook()
    ws = wb.active; ws.title = "Por vencer (7 días)"
    ws["A1"] = "PLANES POR VENCER — PRÓXIMOS 7 DÍAS"; ws["A1"].font = Font(name=A, bold=True, size=13, color="C00000")
    ws["A2"] = f"Generado {data.get('generatedAt','')} · hoy {data.get('today','')}"
    ws["A2"].font = Font(name=A, size=9, color="808080")
    H = ["Vence el", "Días", "Atleta", "Plan"]
    r0 = 4
    for i, h in enumerate(H, 1):
        c = ws.cell(row=r0, column=i, value=h); c.fill = HEAD_FILL; c.font = HEAD_FONT; c.border = BORDER
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    widths(ws, [13, 7, 24, 48])
    rr = r0 + 1
    for p in data.get("expiring", []):
        days = p.get("days_left")
        for i, v in enumerate([p.get("last_session_date", ""), days, p.get("athlete_name", ""), p.get("title", "")], 1):
            cell = ws.cell(row=rr, column=i, value=v); cell.font = CELL; cell.border = BORDER
            cell.alignment = CTR if i in (1, 2) else TOP
        if days is not None and days <= 2:
            for i in range(1, 5): ws.cell(row=rr, column=i).fill = F_HOT
            ws.cell(row=rr, column=2).font = HOT
        elif days is not None and days <= 4:
            for i in range(1, 5): ws.cell(row=rr, column=i).fill = F_PAIN
        rr += 1
    if not data.get("expiring"):
        ws.cell(row=rr, column=1, value="Sin planes por vencer esta semana").font = Font(name=A, size=10, italic=True, color="808080")

    we = wb.create_sheet("Vencidos sin renovar")
    H2 = ["Venció el", "Hace (días)", "Atleta", "Plan"]
    we.append(H2); header(we, len(H2)); widths(we, [13, 11, 24, 48])
    for p in data.get("expired", []):
        we.append([p.get("last_session_date", ""), p.get("days_ago"), p.get("athlete_name", ""), p.get("title", "")])
        i = we.max_row
        for k in range(1, 5):
            cell = we.cell(row=i, column=k); cell.font = CELL; cell.border = BORDER
            cell.alignment = CTR if k in (1, 2) else TOP
        for k in range(1, 5): we.cell(row=i, column=k).fill = F_HOT

    if data.get("competitions"):
        wcp = wb.create_sheet("Competencias próximas")
        H3 = ["Fecha", "Atleta", "Competición"]
        wcp.append(H3); header(wcp, len(H3)); widths(wcp, [13, 24, 44])
        for c in data["competitions"]:
            wcp.append([c.get("date", ""), c.get("athlete_name", ""), c.get("name", "")])
            i = wcp.max_row
            for k in range(1, 4):
                cell = wcp.cell(row=i, column=k); cell.font = CELL; cell.border = BORDER
                cell.alignment = CTR if k == 1 else TOP

    wb.save(OUT)
    for sheet in ("Vencidos sin renovar",):
        drop_ghost_row(OUT, sheet)
    print(f"saved {OUT} | por vencer {len(data.get('expiring', []))} | vencidos {len(data.get('expired', []))} | competencias {len(data.get('competitions', []))}")


# ---------------------------------------------------------------------------
# Modo dossier (histórico de un plan por atleta)
# ---------------------------------------------------------------------------

def build_dossier(data_file):
    data = json.load(open(data_file))
    plan = data.get("plan", {}); sess = data.get("sessions", []); comps = data.get("competitions", [])
    athlete = data.get("athleteName") or (data.get("athlete") or {}).get("name") or f"plan-{plan.get('id')}"
    OUT = f"report/Dossier-{slug(athlete)}.xlsx"
    wb = Workbook()
    ws = wb.active; ws.title = "Sesiones"
    ws["A1"] = f"DOSSIER — {athlete}"; ws["A1"].font = Font(name=A, bold=True, size=14, color="F15A24")
    ws["A2"] = f"Plan: {plan.get('title','')} · {plan.get('startDate','')} → {plan.get('endDate','')} · generado {data.get('generatedAt','')}"
    ws["A2"].font = Font(name=A, size=9, color="808080")
    ct = data.get("counts", {})
    ws["A3"] = f"Sesiones reportadas: {ct.get('sessions', len(sess))} · sRPE≥6: {ct.get('hot','-')} · Dolor/molestia: {ct.get('pain','-')} · Competencias: {ct.get('competitions', len(comps))}"
    ws["A3"].font = Font(name=A, size=10, bold=True)
    H = ["Fecha", "Mom.", "Sesión", "Deporte", "Tipo", "Clave", "RPE prog.", "sRPE", "Dur. real", "Dist. real", "Feedback / comentario", "Dolor", "Estado"]
    r0 = 5
    for i, h in enumerate(H, 1):
        c = ws.cell(row=r0, column=i, value=h); c.fill = HEAD_FILL; c.font = HEAD_FONT; c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 26
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    widths(ws, [12, 6, 26, 11, 18, 7, 9, 7, 10, 11, 56, 8, 11])
    rr = r0 + 1
    for s in sess:
        hot = (s.get("rpe") or 0) >= 6; pain = s.get("pain")
        vals = [s.get("date", ""), s.get("moment", ""), s.get("title", ""), s.get("sport", ""),
                s.get("sessionType", ""), "★" if s.get("keySession") else "", s.get("progRpe"), s.get("rpe"),
                s.get("duration", "") or "—", s.get("distance", "") or "—", (s.get("comment") or "").strip(),
                "Sí" if pain else "—", s.get("statusDisplay") or s.get("status", "")]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=rr, column=i, value=v); cell.font = CELL; cell.border = BORDER
            cell.alignment = WRAP if i in (3, 11) else (CTR if i in (1, 2, 6, 7, 8, 9, 10, 12) else TOP)
        rf = F_PAIN if pain else (F_HOT if hot else None)
        if rf:
            for i in range(1, len(H) + 1):
                if i != 5: ws.cell(row=rr, column=i).fill = rf
        tipo_cell(ws.cell(row=rr, column=5), s.get("intensity"))
        sc = ws.cell(row=rr, column=8); sc.font = HOT if hot else BOLD
        if pain: ws.cell(row=rr, column=12).font = Font(name=A, size=10, bold=True, color="C00000")
        rr += 1

    if comps:
        wc = wb.create_sheet("Competencias")
        H2 = ["Fecha", "Competición", "Prio", "Estado", "sRPE", "Dur.", "Dist.", "Feedback"]
        wc.append(H2); header(wc, len(H2)); widths(wc, [12, 32, 6, 12, 7, 10, 11, 46])
        for c in comps:
            wc.append([c.get("date", ""), c.get("name", ""), c.get("priority", ""), c.get("status", ""),
                       c.get("rpe"), c.get("duration", "") or "—", c.get("distance", "") or "—", c.get("comment", "")])
            i = wc.max_row
            for k in range(1, len(H2) + 1):
                cell = wc.cell(row=i, column=k); cell.font = CELL; cell.border = BORDER
                cell.alignment = WRAP if k == 8 else (CTR if k in (1, 3, 5, 6, 7) else TOP)
            wc.cell(row=i, column=5).font = HOT if (c.get("rpe") or 0) >= 6 else BOLD
    wb.save(OUT)
    if comps:
        drop_ghost_row(OUT, "Competencias")
    print(f"saved {OUT} | sesiones {len(sess)} | competencias {len(comps)}")
    return OUT


if __name__ == "__main__":
    mode = sys.argv[1] if len(sys.argv) > 1 else "daily"
    if mode == "expiry":
        build_expiry()
    elif mode == "dossier":
        build_dossier(sys.argv[2])
    else:
        build_daily()
