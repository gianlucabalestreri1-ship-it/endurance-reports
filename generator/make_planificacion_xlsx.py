#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Planilla de PLANIFICACIÓN por atleta — réplica del archivo modelo del coach.
3 pestañas: Ficha (perfil/intake + HISTORIAL DEPORTIVO) · Rendimiento (PRs por año +
registro de competencias desde TODOS los planes de ET) · Planificación (grilla de
periodización con código de color del manual: COMPETENCIA violeta, TEST verde,
HIT rojo, ThT amarillo, LIT azul). Lee un JSON {name,email,ficha,historial,competitions}
y escribe report/Planificacion-<name>.xlsx. Se sube a Drive convertido a Google Sheets.
Sin fórmulas. Campos que ET no tiene quedan vacíos (los carga el coach)."""
import json, sys, re
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

data = json.load(open(sys.argv[1])) if len(sys.argv) > 1 else {"name":"Atleta","email":"","ficha":{}}
name = data.get("name","Atleta")
ficha = data.get("ficha",{})
historial = data.get("historial",{})
comps = data.get("competitions",[])

# paleta
PURPLE="674EA7"; MAUVE="C27BA0"; GREEN="6AA84F"; RED="E06666"; YELLOW="FFD966"; BLUE="6FA8DC"
DARK="1F3A5F"; LABF="EEF2F7"; INK="1A1F2B"
def fill(c): return PatternFill("solid", fgColor=c)
WHITEF=Font(color="FFFFFF", bold=True, size=11)
INKF=Font(color=INK, size=11)
BOLD=Font(bold=True, color=INK, size=11)
thin=Side(style="thin", color="B8C0CC"); BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
CEN=Alignment(horizontal="center", vertical="center", wrap_text=True)
CEN1=Alignment(horizontal="center", vertical="center", wrap_text=False)   # sin ajuste (rótulos de 1 línea)
LEFT=Alignment(horizontal="left", vertical="center", wrap_text=True)

def dmy(s):
    try: return date.fromisoformat(s).strftime("%d/%m/%Y")
    except Exception: return s or ""

wb=Workbook()

# ================= PESTAÑA 1: FICHA =================
ws=wb.active; ws.title="Ficha"
ws.column_dimensions["A"].width=34; ws.column_dimensions["B"].width=72
CAMPOS=["NOMBRE Y APELLIDO","FECHA DE NACIMIENTO","SEXO","LUGAR DE RESIDENCIA",
"NÚMERO DE TELÉFONO","EMAIL","CONDICIÓN MÉDICA","LESIONES","MEDICACIÓN","CIRUGÍAS",
"FUMA/ALCOHOL","DIETA","HORAS DE SUEÑO","TRABAJO","EXPERIENCIA EN RUNNING",
"COMPETENCIAS QUE HA PARTICIPADO","EXPERIENCIA CON ENTRENADOR",
"¿PRACTICA OTRAS DISCIPLINAS A LA PAR?","MOTIVACIONES PARA ENTRENAR",
"DÍAS A LA SEMANA DISPONIBLES","CUANDO NO PUEDE ENTRENAR"]
LARGAS={"CONDICIÓN MÉDICA","DIETA","COMPETENCIAS QUE HA PARTICIPADO","MOTIVACIONES PARA ENTRENAR"}
r=1
for campo in CAMPOS:
    lc=ws.cell(r,1,campo); lc.fill=fill(LABF); lc.font=BOLD; lc.alignment=LEFT; lc.border=BORD
    vc=ws.cell(r,2,ficha.get(campo,"")); vc.font=INKF; vc.alignment=LEFT; vc.border=BORD
    if campo in LARGAS: ws.row_dimensions[r].height=58
    r+=1
# --- HISTORIAL DEPORTIVO ---
r+=1
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=2)
hc=ws.cell(r,1,"HISTORIAL DEPORTIVO"); hc.fill=fill(DARK); hc.font=WHITEF; hc.alignment=Alignment(horizontal="left",vertical="center")
ws.cell(r,2).border=BORD; hc.border=BORD; r+=1
HIST=["Años Corriendo","Experiencia Previa (deportes)","Última Carrera y Ritmo",
"Volumen Promedio Semanal (km)","Entrenamientos Semanales Actuales (#)","KMS PICO (máx. semanal histórico)"]
for campo in HIST:
    lc=ws.cell(r,1,campo); lc.fill=fill(LABF); lc.font=BOLD; lc.alignment=LEFT; lc.border=BORD
    vc=ws.cell(r,2,historial.get(campo,"")); vc.font=INKF; vc.alignment=LEFT; vc.border=BORD
    r+=1

# ================= PESTAÑA 2: RENDIMIENTO =================
ws2=wb.create_sheet("Rendimiento")
cols2=["Año","Mejor 5K","Mejor 10K","Mejor 21K","Mejor 42K","Test de Milla",
"VO2 Max estimado","FCM","VT1","VT2","KMS PICO"]
for i,w in enumerate([12,12,12,12,12,13,14,10,10,10,12],1):
    ws2.column_dimensions[get_column_letter(i)].width=w
for i,h in enumerate(cols2,1):
    c=ws2.cell(1,i,h); c.fill=fill(PURPLE); c.font=WHITEF; c.alignment=CEN; c.border=BORD
for rr in range(2,7):
    for i in range(1,len(cols2)+1):
        c=ws2.cell(rr,i,""); c.border=BORD; c.alignment=CEN
# ---- registro de competencias (desde TODOS los planes ET) ----
tr=8
ws2.merge_cells(start_row=tr,start_column=1,end_row=tr,end_column=6)
tc=ws2.cell(tr,1,"Registro de competencias (histórico Endurance Tool)")
tc.font=BOLD; tc.alignment=Alignment(horizontal="left",vertical="center"); tc.fill=fill(LABF)
for i in range(1,7): ws2.cell(tr,i).border=BORD
tr+=1
cols2b=["Fecha","Tiempo (mm:ss)","Ritmo (min/km)","Frecuencia cardíaca (med/max)","Condiciones","Carrera / Notas"]
for i,h in enumerate(cols2b,1):
    c=ws2.cell(tr,i,h); c.fill=fill(DARK); c.font=WHITEF; c.alignment=CEN; c.border=BORD
tr+=1
nrows=max(12, len(comps)+2)
for k in range(nrows):
    c=comps[k] if k < len(comps) else None
    fecha = dmy(c["date"]) if c else ""
    nota  = (f'{c.get("name","")}' + (f' · prioridad {c.get("priority")}' if c and c.get("priority") else '')) if c else ""
    if c and c.get("rpe") is not None: nota += f' · sRPE {c["rpe"]}'
    if c and c.get("comment"): nota += f' · {c["comment"]}'
    vals=[fecha,"","","","",nota]
    for i,v in enumerate(vals,1):
        cc=ws2.cell(tr,i,v); cc.border=BORD
        cc.alignment=(LEFT if i==6 else CEN); cc.font=INKF
    tr+=1

# ================= PESTAÑA 3: PLANIFICACIÓN =================
ws3=wb.create_sheet("Planificación")
NMC=28
ws3.column_dimensions["A"].width=19
for i in range(2,2+NMC): ws3.column_dimensions[get_column_letter(i)].width=4.6
# título (combinado para que no se apile vertical)
ws3.merge_cells(start_row=1,start_column=1,end_row=1,end_column=14)
tc=ws3.cell(1,1,"PLANTILLA PLANIFICACIÓN DEL ENTRENAMIENTO")
tc.font=Font(bold=True,size=13,color=INK); tc.alignment=Alignment(horizontal="left",vertical="center")
# objetivos (los carga el coach - NO inventar)
ws3.cell(3,1,"Objetivos").font=BOLD
for lab,row in [("Carrera objetivo:",4),("Objetivos generales:",5),("Objetivos específicos:",6),("Objetivos operativos:",7)]:
    c=ws3.cell(row,1,lab); c.font=Font(bold=True,size=10,color=INK); c.alignment=Alignment(horizontal="left",vertical="center")
r=9
def grid_label(row, text, bg=None, fg=INK):
    c=ws3.cell(row,1,text); c.font=Font(bold=True,size=10,color=fg); c.alignment=CEN1; c.border=BORD
    c.fill=fill(bg if bg else LABF); return c
def grid_cells(row, bg=None, vals=None):
    darkbg = bg in (PURPLE,MAUVE,GREEN,RED,DARK)
    for j in range(NMC):
        v=(vals[j] if vals and j<len(vals) else "")
        c=ws3.cell(row,2+j,v); c.alignment=CEN; c.border=BORD
        c.font=Font(size=8, color=("FFFFFF" if darkbg else INK))
        if bg: c.fill=fill(bg)
grid_label(r,"PERÍODO", PURPLE, "FFFFFF"); grid_cells(r, PURPLE); r+=1
grid_label(r,"MESOCICLO"); grid_cells(r); r+=1
grid_label(r,"OBJETIVO (END/LT/VO2/RP)"); grid_cells(r); ws3.row_dimensions[r].height=26; r+=1
grid_label(r,"MICROCICLO"); grid_cells(r, vals=[str(i) for i in range(1,NMC+1)]); r+=1
grid_label(r,"KILÓMETROS"); grid_cells(r); r+=1
METHODS=[("COMPETENCIA",MAUVE,"FFFFFF"),("TEST",GREEN,"FFFFFF"),
("SPRINT",RED,"FFFFFF"),("GLUCOLÍTICA",RED,"FFFFFF"),("VO2MAX",RED,"FFFFFF"),
("SUPRAUMBRAL",RED,"FFFFFF"),("UMBRAL",YELLOW,INK),("TEMPO",YELLOW,INK),
("ENDURANCE",BLUE,INK),("RECUPERACIÓN",BLUE,INK)]
for lab,bg,fg in METHODS:
    grid_label(r,lab,bg,fg); grid_cells(r,bg); r+=1
grid_label(r,"COMENTARIOS"); grid_cells(r); ws3.row_dimensions[r].height=30; r+=1

out=f"report/Planificacion-{re.sub(r'[^A-Za-z0-9 ]','',name).strip()}.xlsx"
wb.save(out); print("OK ->", out)
