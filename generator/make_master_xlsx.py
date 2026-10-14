# -*- coding: utf-8 -*-
"""Generador del MASTER por atleta (planilla integral).
Uso: python3 generator/make_master_xlsx.py report/master-<id>-data.json
Lee ese JSON + data/*.json (referencias) y escribe report/Master-<Atleta>.xlsx
Hojas: PANEL · DASHBOARD · PLANIFICACIÓN ACTUAL · HISTORIAL · PERFILADO · RECOMENDACIÓN.
Datos sensibles -> entregar como archivo, no publicar."""
import json, sys, re, datetime
from openpyxl import load_workbook, Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = "/home/user/endurance-reports/"
D = json.load(open(sys.argv[1]))
def ref(n): return json.load(open(ROOT + "data/" + n + ".json"))
ZB = ref("modelo-zonas-base"); DIST = ref("distribucion-carga-periodo")
MET = ref("metodos-entrenamiento"); PROG = ref("progresion-volumen-anual")
BAR = ref("baremo-perfil-fondo"); JDC = ref("sesiones-microciclos-jdc")
LIBE = ref("libreria-et")
TPL = ROOT + "plantilla.xlsx"
TODAY = D["generatedAt"]
RO = D["rowOrder"]; MBLK = 5

INK = "1A1F2B"
HEAD = PatternFill("solid", fgColor="1F3A5F"); HF = Font(color="FFFFFF", bold=True, size=9)
PLANF = PatternFill("solid", fgColor="2E4A6B"); LABF = PatternFill("solid", fgColor="EEF2F7")
COACH = PatternFill("solid", fgColor="FFF6E6"); SUBF = Font(italic=True, size=9, color="5B6472")
TITF = Font(bold=True, size=16, color=INK); LF = Font(bold=True, size=9, color="1F3A5F"); CF = Font(size=9)
thin = Side(style="thin", color="C9D2DF"); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
hoyS = Side(style="medium", color="C0392B")
C = Alignment("center", vertical="center", wrap_text=True); L = Alignment("left", vertical="center", wrap_text=True)
COL = {"SPRINT": ("D64550", "F4CCCC"), "GLUCOLÍTICA": ("D64550", "F4CCCC"), "VO2MAX": ("D64550", "F4CCCC"),
       "SUPRAUMBRAL": ("E69138", "FCE5CD"), "UMBRAL": ("E8B800", "FFF2CC"), "TEMPO": ("E8B800", "FFF2CC"),
       "ENDURANCE": ("4A86C7", "CFE2F3"), "RECUPERACIÓN": ("4A86C7", "CFE2F3"), "FUERZA": ("999999", "D9D9D9")}
LBL = {"SPRINT": "Sprint / Neuromuscular", "GLUCOLÍTICA": "Glucolítico", "VO2MAX": "VO2Max",
       "SUPRAUMBRAL": "Supraumbral", "UMBRAL": "Umbral", "TEMPO": "Tempo",
       "ENDURANCE": "Endurance / Fondo", "RECUPERACIÓN": "Recuperación", "FUERZA": "Fuerza"}
PHF = {"PREPARATORIO": PatternFill("solid", fgColor="D9D2E9"), "ESPECÍFICO": PatternFill("solid", fgColor="D0E0E3"),
       "COMPETITIVO": PatternFill("solid", fgColor="FCE5CD"), "MIXTO": PatternFill("solid", fgColor="EFEFEF")}

def fd(s):
    if not s: return ""
    p = str(s).split("-"); return f"{p[2]}/{p[1]}/{p[0]}" if len(p) == 3 else str(s)
def fdm(s):
    if not s: return ""
    p = str(s).split("-"); return f"{p[2]}/{p[1]}" if len(p) == 3 else str(s)
def mcell(ws, r, c1, c2, val, fill=None, font=None, align=C):
    if c2 > c1: ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    cell = ws.cell(r, c1, val)
    if fill: cell.fill = fill
    if font: cell.font = font
    cell.alignment = align
    for cc in range(c1, c2 + 1): ws.cell(r, cc).border = BORD
    return cell
def spans(weeks):
    out = []; i = 0; n = len(weeks)
    while i < n:
        mn = weeks[i]["meso"]; j = i
        while j + 1 < n and weeks[j + 1]["meso"] == mn: j += 1
        out.append((mn, i, j)); i = j + 1
    return out
def pace_to_sec(p):  # "5'11\"/km" o "4'37\" / km" -> seg/km
    m = re.search(r"(\d+)'(\d+)", p or "")
    return int(m.group(1)) * 60 + int(m.group(2)) if m else None
def time_to_sec(t):
    if not t: return None
    t = t.replace("h", ":").replace("'", ":").strip()
    parts = [x for x in re.split(r"[:\s]+", t) if x.isdigit()]
    if len(parts) == 3: return int(parts[0]) * 3600 + int(parts[1]) * 60 + int(parts[2])
    if len(parts) == 2: return int(parts[0]) * 60 + int(parts[1])
    return None

# ---------- auto-síntesis de feedback por meso ----------
def synth(m):
    cs = m.get("comments", [])
    if not cs: return ""
    n = len(cs)
    extra = (" · sRPE~" + m["srpeAvg"]) if m.get("srpeAvg") else ""
    pain = (" · ⚠ " + str(len(m["pains"])) + " molestia(s)") if m.get("pains") else ""
    sal = [c["c"] for c in cs if (c.get("rpe") or 0) >= 6][:1] or [c["c"] for c in cs][:2]
    return f"{n} feedbacks{extra}{pain}. Ej.: " + " / ".join(x[:70] for x in sal) + "  (auto — editar)"

# ---------- carreras reales (con FIT) ----------
RACES = [r for r in D.get("races", []) if r.get("pace") and r.get("km")]
def race_dist(name):
    n = name or ""
    for d in ["42", "21", "15", "10", "5"]:
        if d in n: return d + "K"
    return ""
def best_pr(dist_tag):
    cand = [r for r in RACES if race_dist(r["name"]) == dist_tag and r.get("time")]
    return min(cand, key=lambda r: time_to_sec(r["time"]) or 9e9) if cand else None

# ---------- baremo: nivel por vVO2max ----------
def vvo2_kmh():
    z = D.get("vVO2maxZone", "")
    paces = re.findall(r"(\d+)'(\d+)", z)
    if len(paces) >= 2:
        s = [int(a) * 60 + int(b) for a, b in paces[:2]]; mid = sum(s) / 2
        return 3600.0 / mid
    return None
def nivel_baremo():
    v = vvo2_kmh()
    if not v: return None
    col = "vvo2max_m" if D["profile"].get("gender") == "Femenino" else "vvo2max_h"
    cl = None
    for row in sorted(BAR["rows"], key=lambda r: r["aiym"]):
        if row.get(col, 0) <= v: cl = row
    return (cl, v)

PROF = D.get("profile", {})
ADV = MET["adaptacion_objetivo_por_periodo"]

# --- recomputar molestias ACTIVAS (recuperación/menstrual NO cuentan) ---
_PAIN = re.compile(r"(dolor|molest|duele|duelen|lesi[oó]n|pinchazo|contractur|sobrecarg|tir[oó]n|calambr|inflama|resent|pincha|hinch|adolor)", re.I)
_NEG = re.compile(r"\b(sin|no|ni|ning\w*|cero|nunca|tampoco)\b", re.I)
_RECOV = re.compile(r"(recuper|ya pas[oó]|ya est[aá]|super[eé]|mejor|sin problema|casi sin|se fue|remit|ya no|molest[oó] menos|cada vez menos)", re.I)
_MENS = re.compile(r"(menstrua|menstr|la regla|per[ií]odo|ciclo)", re.I)
def _active_pain(t):
    t = (t or "").lower()
    for m in _PAIN.finditer(t):
        i = m.start(); clause = re.split(r"[.;,:!?]|\bpero\b|\baunque\b|\bsalvo\b", t[max(0, i - 60):i])[-1]; win = t[max(0, i - 60):i + 40]
        if _NEG.search(clause): continue
        if _RECOV.search(win): continue
        if _MENS.search(win): continue
        return True
    return False
for _p in D["plans"]:
    for _m in _p["mesos"]:
        _m["pains"] = [f'{c["date"]}: {c["c"][:70]}' for c in _m.get("comments", []) if _active_pain(c["c"])]

wb = load_workbook(TPL)
# ============ se sigue en parte 2 (anexada por el script generador) ============

# =================== helpers de render ===================
def anal_for(start, end):
    return next((a["notes"] for a in D.get("analisis", []) if start <= a["date"] <= end), "")
def pace_mid(zone):
    ps = re.findall(r"(\d+)'(\d+)", zone or "")
    if len(ps) >= 2:
        s = [int(a) * 60 + int(b) for a, b in ps[:2]]; return sum(s) / 2
    return None
def fmt_pace(sec):
    if not sec: return ""
    return f"{int(sec // 60)}'{int(round(sec % 60)):02d}\"/km"

# meses dinámicos (2 atrás, actual, 3 adelante)
_cy, _cm = int(TODAY[:4]), int(TODAY[5:7])
MONTHS = []
for off in range(-2, 4):
    mm = _cm + off; yy = _cy + (mm - 1) // 12; mm = (mm - 1) % 12 + 1; MONTHS.append((yy, mm))
_MN = ["", "ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]
def mk(y, m): return f"{y}-{m:02d}"
def prevm(y, m): return (y - 1, 12) if m == 1 else (y, m - 1)
def month_comments(y, m):
    k = mk(y, m); cs = []
    for p in D["plans"]:
        for mm in p["mesos"]:
            for c in mm.get("comments", []):
                if (c["date"] or "")[:7] == k: cs.append(c)
    return cs
def month_synth(y, m):
    cs = month_comments(y, m)
    if not cs: return ""
    sal = [c["c"] for c in cs if (c.get("rpe") or 0) >= 6][:1] or [c["c"] for c in cs][:1]
    return f"{len(cs)} feedbacks. " + " / ".join(x[:60] for x in sal) + " (auto)"
RACEM = {}
for r in RACES:
    RACEM.setdefault((r["date"] or "")[:7], []).append(f'{r["name"]} {r["pace"]}/FC{r["hr"]}')

# =================== DASHBOARD ===================
ws = wb["DASHBOARD"]
for a in ["B5","B6","B7","B8","B9","B10","B11","I5","I19","I20","I21","I22","I23","I24","I25","I26","I29","I30","B14","B15","B16"]:
    try: ws[a] = None
    except Exception: pass
ws["B4"] = D["athlete"]
ws["B5"] = f'{PROF.get("age","")} - {PROF.get("gender","")}'.strip(" -")
ws["B7"] = PROF.get("province", "")
ws["B8"] = PROF.get("height", "")
ws["B9"] = PROF.get("weight", "")
ws["I4"] = D["anios"]["texto"]
# última carrera
if RACES:
    lastr = max(RACES, key=lambda r: r["date"])
    ws["I6"] = f'{lastr["name"]} ({fd(lastr["date"])}) · {lastr["pace"]} · FC {lastr["hr"]}'
ws["G7"] = "Volumen semanal (máx hist / prom mes)"
if D.get("volMax"):
    ws["I7"] = f'Máx histórico: {D["volMax"]["km"]} km ({fd(D["volMax"]["date"])}) · Prom. últ. mes: {D["volAvg"]} km/sem ({"/".join(str(x) for x in D["lastWeeksKm"])})'
ws["I7"].alignment = L
# MOTOR
vv_sec = pace_mid(D.get("vVO2maxZone", ""))
ee = min([r for r in RACES if time_to_sec(r.get("time")) and abs(time_to_sec(r["time"]) - 3600) < 1800], key=lambda r: abs(time_to_sec(r["time"]) - 3600), default=None)
ul = min([r for r in RACES if time_to_sec(r.get("time")) and abs(time_to_sec(r["time"]) - 2400) < 1200], key=lambda r: abs(time_to_sec(r["time"]) - 2400), default=None)
if vv_sec: ws["B15"] = fmt_pace(vv_sec)
ws["B16"] = (f'{ee["pace"]} — {ee["name"]} ({fd(ee["date"])})' if ee else "s/dato (ref: carrera ~1h)")
ws["A17"] = "Umbral de lactato"; ws["B17"] = (f'{ul["pace"]} — {ul["name"]} ({fd(ul["date"])})' if ul else "s/dato (ref: carrera ~40' / 8-10K)")
if vv_sec and ee and pace_to_sec(ee["pace"]):
    pct = round((pace_to_sec(ee["pace"]) - vv_sec) / vv_sec * 100)
    clasif = "POTENTE (FT-accesible)" if pct > 15 else ("EQUILIBRADO" if pct >= 10 else "EFICIENTE (ST-dominante)")
    ws["B14"] = f"≈ {pct}% — {clasif}"
from copy import copy as _copy
ws["A17"].fill = _copy(ws["A16"].fill); ws["A17"].border = _copy(ws["A16"].border); ws["A17"].font = _copy(ws["A16"].font)
ws["B17"].border = _copy(ws["B16"].border)
_mid = Alignment("center", vertical="center", wrap_text=True)
ws["A13"].alignment = _mid
for rr in range(14, 18):
    for co in ("A", "B"):
        c = ws[f"{co}{rr}"]; c.alignment = _mid
        if not (c.border and c.border.left and c.border.left.style): c.border = BORD

# =================== PLANIFICACIÓN (análisis por meso + mapa) ===================
TROWS = [("Período","a"),("Competencias","a"),("Tipo / Fase (periodización)","c"),("Volumen semanal objetivo","c"),
         ("Volumen real / adherencia","a"),("Intensidad prioritaria","c"),("Tipos de entrenamiento","c"),
         ("Sesiones de referencia (mes)","a"),("Valoraciones / tests","a"),("Comentarios del atleta (mes)","a"),
         ("Molestias (del feedback)","a"),("Obs / análisis del entrenador","c"),("Taper","c")]
RB = {r["name"]: r for r in RACES}
def intens_sug(z):
    if z["mid"] >= 50: return f'Umbral/Tempo (mid {z["mid"]}%)'
    if z["low"] >= 50: return f'Base Z1-2 (low {z["low"]}%)'
    if z["high"] >= 25: return f'VO2/Alta (high {z["high"]}%)'
    return f'Mixto {z["high"]}/{z["mid"]}/{z["low"]}'
def comp_txt(m):
    out = []
    for c in m["comps"]:
        r = RB.get(c["name"], {})
        ex = (f' {r["pace"]} FC{r["hr"]}' if r.get("pace") else "")
        out.append(f'{fdm(c["date"])} {c["name"]} [{c["prio"]}]{ex}')
    return " / ".join(out)
def refs_txt(m):
    return "\n".join(f'{c["name"]}: {RB[c["name"]]["pace"]} · FC {RB[c["name"]]["hr"]}' for c in m["comps"] if c["name"] in RB and RB[c["name"]].get("pace"))
def tval(label, m):
    z = m["zonesPct"]
    return {"Período": f'{fdm(m["start"])}–{fdm(m["end"])}', "Competencias": comp_txt(m),
            "Tipo / Fase (periodización)": m["phaseInf"], "Volumen semanal objetivo": f'~{m["volProg"]} (ET prog.)',
            "Volumen real / adherencia": f'{m["executed"]}/{m["sessions"]} ses · sRPE {m["srpeAvg"] or "—"}',
            "Intensidad prioritaria": intens_sug(z), "Tipos de entrenamiento": "",
            "Sesiones de referencia (mes)": refs_txt(m),
            "Valoraciones / tests": "\n".join(m["tests"]) if m["tests"] else "",
            "Comentarios del atleta (mes)": synth(m),
            "Molestias (del feedback)": ("⚠ " + "\n".join(m["pains"])) if m["pains"] else "",
            "Obs / análisis del entrenador": anal_for(m["start"], m["end"]), "Taper": ""}[label]
def render_table(ws, p, r):
    mesos = p["mesos"]; nm = len(mesos); last = 1 + nm * MBLK
    mcell(ws, r, 1, last, "ANÁLISIS POR MESOCICLO  (amarillo = lo definís vos · claro = del asistente)", PatternFill("solid", fgColor="34495E"), Font(bold=True, size=10, color="FFFFFF"), L); r += 1
    ws.cell(r, 1, "").fill = HEAD; ws.cell(r, 1).border = BORD
    for i, m in enumerate(mesos): mcell(ws, r, 2 + i * MBLK, 1 + (i + 1) * MBLK, f'Meso {m["number"]}', HEAD, HF)
    r += 1
    BIG = {"Comentarios del atleta (mes)", "Obs / análisis del entrenador"}
    for label, owner in TROWS:
        ws.cell(r, 1, label).fill = (COACH if owner == "c" else LABF); ws.cell(r, 1).font = LF; ws.cell(r, 1).alignment = L; ws.cell(r, 1).border = BORD
        for i, m in enumerate(mesos):
            al = L if label in BIG or label in ("Competencias", "Sesiones de referencia (mes)", "Molestias (del feedback)", "Valoraciones / tests", "Tipos de entrenamiento") else C
            mcell(ws, r, 2 + i * MBLK, 1 + (i + 1) * MBLK, tval(label, m), COACH if owner == "c" else None, CF, al)
        if label in BIG: ws.row_dimensions[r].height = 85
        r += 1
    return r + 1
def render_map(ws, p, r):
    weeks = p["weeks"]; n = len(weeks)
    if n == 0: return r
    last = 1 + n
    def band(label, fn, fillfn=None):
        nonlocal r
        ws.cell(r, 1, label).fill = HEAD; ws.cell(r, 1).font = HF; ws.cell(r, 1).alignment = L; ws.cell(r, 1).border = BORD
        i = 0
        while i < n:
            v = fn(weeks[i]); j = i
            while j + 1 < n and fn(weeks[j + 1]) == v and v not in (None, ""): j += 1
            mcell(ws, r, 2 + i, 2 + j, v if v else "", fillfn(v) if (fillfn and v) else None, Font(bold=True, size=9)); i = j + 1
        r += 1
    mcell(ws, r, 1, last, "MAPA DE PERIODIZACIÓN  (fuerte=ejecutado · pálido=pendiente · borde rojo=HOY)", PatternFill("solid", fgColor="34495E"), Font(bold=True, size=10, color="FFFFFF"), L); r += 1
    _FL = {"PREPARATORIO": "Base", "ESPECÍFICO": "Específico", "COMPETITIVO": "Competitivo", "MIXTO": "Mixto"}
    band("PERÍODO", lambda w: _FL.get(w["phase"], w["phase"]), lambda v: PHF.get({"Base":"PREPARATORIO","Específico":"ESPECÍFICO","Competitivo":"COMPETITIVO"}.get(v, "MIXTO"), PHF["MIXTO"]))
    band("MESOCICLO", lambda w: f'Meso {w["meso"]}')
    ws.cell(r, 1, "KILÓMETROS (sem)").fill = PatternFill("solid", fgColor="223B57"); ws.cell(r, 1).font = HF; ws.cell(r, 1).alignment = L; ws.cell(r, 1).border = BORD
    for i, w in enumerate(weeks):
        cc = ws.cell(r, 2 + i, w["km"] or ""); cc.alignment = C; cc.font = Font(size=8, bold=True); cc.border = BORD
    r += 1
    ws.cell(r, 1, "MICROCICLO").fill = HEAD; ws.cell(r, 1).font = HF; ws.cell(r, 1).alignment = L; ws.cell(r, 1).border = BORD
    for i, w in enumerate(weeks):
        dd = str(w["start"]).split("-"); lab = f'{w["n"]}\n{dd[2]}/{dd[1]}' if len(dd) == 3 else str(w["n"])
        cc = ws.cell(r, 2 + i, lab); cc.alignment = C; cc.font = Font(size=8, bold=True); cc.border = BORD
        if w["start"] <= TODAY <= w["end"]: cc.fill = PatternFill("solid", fgColor="C0392B"); cc.font = Font(size=8, bold=True, color="FFFFFF")
    r += 1
    ws.cell(r, 1, "COMPETENCIA").fill = PatternFill("solid", fgColor="B07AA1"); ws.cell(r, 1).font = Font(bold=True, size=9, color="FFFFFF"); ws.cell(r, 1).alignment = L; ws.cell(r, 1).border = BORD
    for i, w in enumerate(weeks):
        cc = ws.cell(r, 2 + i, "🏁" if w["comp"] else ""); cc.alignment = C; cc.border = BORD
        if w["comp"]: cc.fill = PatternFill("solid", fgColor="C27BA0"); cc.font = Font(size=9, color="FFFFFF"); cc.comment = Comment(w["comp"], "ET")
    r += 1
    ws.cell(r, 1, "TEST").fill = PatternFill("solid", fgColor="6AA84F"); ws.cell(r, 1).font = Font(bold=True, size=9, color="FFFFFF"); ws.cell(r, 1).alignment = L; ws.cell(r, 1).border = BORD
    for i, w in enumerate(weeks):
        cc = ws.cell(r, 2 + i, "T" if w["test"] else ""); cc.alignment = C; cc.border = BORD
        if w["test"]: cc.fill = PatternFill("solid", fgColor="93C47D"); cc.font = Font(size=9, bold=True)
    r += 1
    for R in RO:
        strong, pale = COL[R]
        ws.cell(r, 1, LBL[R]).fill = LABF; ws.cell(r, 1).font = LF; ws.cell(r, 1).alignment = L; ws.cell(r, 1).border = BORD
        for i, w in enumerate(weeks):
            pe = w["rows"][R]; p_, e_ = pe["p"], pe["e"]
            cc = ws.cell(r, 2 + i, ""); cc.alignment = C; cc.border = BORD
            if p_ > 0:
                if e_ >= p_: cc.fill = PatternFill("solid", fgColor=strong); cc.value = "●" if p_ == 1 else str(e_); cc.font = Font(bold=True, size=9, color="FFFFFF")
                elif e_ > 0: cc.fill = PatternFill("solid", fgColor=strong); cc.value = f"{e_}/{p_}"; cc.font = Font(bold=True, size=8, color="FFFFFF")
                else: cc.fill = PatternFill("solid", fgColor=pale); cc.value = "○" if p_ == 1 else f"·{p_}"; cc.font = Font(size=8, color="7A7A7A")
        r += 1
    return r + 2
def render_plan(ws, p, r):
    last = 1 + max(len(p["mesos"]) * MBLK, len(p["weeks"]))
    mcell(ws, r, 1, last, f'PLAN · {p["title"]}   ·   {fd(p["start"])} → {fd(p["end"])}   ·   {p["status"]}', PLANF, HF, L); r += 1
    r = render_table(ws, p, r)
    r = render_map(ws, p, r)
    return r + 2
def setup(ws, maxw, title, subtitle):
    ws.column_dimensions["A"].width = 26
    for i in range(2, maxw + 2): ws.column_dimensions[get_column_letter(i)].width = 5.7
    mcell(ws, 1, 1, min(maxw + 1, 16), title, None, TITF, L)
    mcell(ws, 2, 1, min(maxw + 1, 16), subtitle, None, SUBF, L)

idx = wb.sheetnames.index("PLANIFICACIÓN"); del wb["PLANIFICACIÓN"]
active = [p for p in D["plans"] if p["status"] == "Active"]; prev = [p for p in D["plans"] if p["status"] != "Active"]
mw = lambda ps: max((max(len(p["mesos"]) * MBLK, len(p["weeks"])) for p in ps), default=18)
wsA = wb.create_sheet("PLANIFICACIÓN ACTUAL", idx)
setup(wsA, mw(active) if active else 18, "PLANIFICACIÓN ACTUAL", "Plan vigente · análisis por mesociclo + mapa de periodización")
r = 4
for p in active: r = render_plan(wsA, p, r)
if not active: wsA.cell(4, 1, "(sin plan activo)")
wsH = wb.create_sheet("HISTORIAL (planes previos)", idx + 1)
setup(wsH, mw(prev) if prev else 18, "HISTORIAL DE PLANIFICACIÓN", "Archivo cronológico")
r = 4
for p in prev: r = render_plan(wsH, p, r)

# =================== PERFILADO ===================
ws = wb["PERFILADO Y COMPETENCIAS"]
for a in ["A93","B93","C93","D93","E93","F93","A104","B104","C104","D104","E104","F104"]:
    try: ws[a] = None
    except Exception: pass
prrow = {"15K": 9, "21K": 10, "42K": 11}
for d, rr in prrow.items():
    rc = best_pr(d)
    if rc:
        ws.cell(rr, 9, rc["time"]); ws.cell(rr, 10, rc["pace"]); ws.cell(rr, 11, rc.get("parciales", ""))
        ws.cell(rr, 15, (str(rc.get("hr", "")) + " ppm") if rc.get("hr") else ""); ws.cell(rr, 23, fd(rc["date"])); ws.cell(rr, 24, rc["name"])
def fill_track(startrow, items):
    r = startrow
    for it in items:
        ws.cell(r, 1, fd(it["date"])); ws.cell(r, 2, it.get("time", "")); ws.cell(r, 3, it.get("pace", ""))
        ws.cell(r, 4, (str(it["hr"]) + " ppm (media)") if it.get("hr") else ""); ws.cell(r, 5, it.get("parciales", "")); ws.cell(r, 6, it["name"] + (f' · {it["km"]}km' if it.get("km") else ""))
        r += 1
fill_track(93, [r for r in RACES if "21" in r["name"]])
fill_track(104, [r for r in RACES if "42" in r["name"]])
fill_track(80, [r for r in RACES if ("10K" in r["name"] or "Sonder" in r["name"]) and "21" not in r["name"]])
r = 125
ws.cell(r, 1, "🏁 COMPETENCIAS (todas — vacío si sin datos)").font = Font(bold=True, size=12, color="1F3A5F"); r += 1
for i, h in enumerate(["Fecha", "Evento", "Prio", "Tiempo", "Ritmo", "FC", "Comentario"], 1):
    c = ws.cell(r, i, h); c.fill = HEAD; c.font = HF; c.border = BORD; c.alignment = C
r += 1
for c in D["allComps"]:
    rc = RB.get(c["name"], {})
    _com = rc.get("comment") or c.get("comment") or ""
    row = [fd(c["date"]), c["name"], c["prio"], rc.get("time", "") or (c.get("dur") or ""), rc.get("pace", ""), (str(rc.get("hr", "")) + " ppm" if rc.get("hr") else ""), _com]
    for i, v in enumerate(row, 1):
        cc = ws.cell(r, i, v); cc.border = BORD; cc.alignment = (L if i in (2, 7) else C); cc.font = CF
    r += 1
r += 2
ws.cell(r, 1, f"⚙️ ZONAS — foto {TODAY} (modelo base Z1-Z8 · valores reales ET)").font = Font(bold=True, size=12, color="1F3A5F"); r += 1
def zblock(title, zones, r):
    ws.cell(r, 1, title).font = LF; r += 1
    for i, h in enumerate(["#", "Zona", "Fase", "Rango"], 1):
        c = ws.cell(r, i, h); c.fill = HEAD; c.font = HF; c.border = BORD; c.alignment = C
    r += 1
    for z in zones:
        for i, v in enumerate([z["n"], z["name"], z["phase"], z["range"]], 1):
            c = ws.cell(r, i, v); c.border = BORD; c.alignment = (L if i in (2, 4) else C); c.font = CF
        r += 1
    return r + 1
if D.get("zones"):
    if D["zones"].get("MAS"): r = zblock("RITMO (MAS) — min/km", D["zones"]["MAS"]["zones"], r)
    if D["zones"].get("HR"): r = zblock("FRECUENCIA CARDÍACA (%máx) — bpm", D["zones"]["HR"]["zones"], r)
ws.cell(r, 1, "Registro de CAMBIOS de zonas (hacia adelante: fecha · criterio · antes→después · test).").font = Font(italic=True, size=9, color="974706")

# borrar hojas no usadas
for nm in ["BIOMECÁNICA", "MOVILIDAD  FUERZA", "ANTROPOMETRÍAS", "ANALÍTICAS SANGUÍNEAS", "INFORME"]:
    if nm in wb.sheetnames: del wb[nm]

# =================== PANEL (capa de decisión) ===================
wsP = wb.create_sheet("PANEL", 0)
act = active[0] if active else (D["plans"][-1] if D["plans"] else None)
weeks = act["weeks"] if act else []
n = len(weeks); LAST = 9  # ancho fijo (A..I): más aire horizontal, menos largo vertical
# columnas que sirven a FUNDAMENTOS, KPIs y al FOCO MENSUAL (meses en filas)
_PW = {"A": 11, "B": 21, "C": 13, "D": 27, "E": 19, "F": 11, "G": 12, "H": 25, "I": 21}
for col, wd in _PW.items(): wsP.column_dimensions[col].width = wd
mbn = {m["number"]: m for m in act["mesos"]} if act else {}
past = [w for w in weeks if w["start"] <= TODAY]
cur = None
for i, w in enumerate(weeks):
    if w["start"] <= TODAY <= w["end"]: cur = (i, w)
if cur is None and past: cur = (len(past) - 1, past[-1])
ci, cw = cur if cur else (0, (weeks[0] if weeks else {"n": 0, "meso": 0, "start": TODAY, "end": TODAY}))
cm = mbn.get(cw["meso"], {})
_pw = [w for w in past if w.get("km")]
vol_act = _pw[-1]["km"] if _pw else (D.get("volAvg") or 0)  # km PROGRAMADO (no depende de feedback)
closed = [w for w in weeks if w["end"] < TODAY]; l4 = closed[-4:]
sp = sum(pe["p"] for w in l4 for k, pe in w["rows"].items() if k != "FUERZA"); se_ = sum(pe["e"] for w in l4 for k, pe in w["rows"].items() if k != "FUERZA")
adh = round(100 * se_ / sp) if sp else 0
# ¿había sesiones programadas pero CERO ejecución registrada? -> el atleta no cargó feedback (no es "entrenó 0")
measurable = sp > 0
no_fb = measurable and se_ == 0
srpes = [m["srpeAvg"] for m in act["mesos"] if m["srpeAvg"]] if act else []
srpe = srpes[-1] if srpes else "—"
cut = (datetime.date.fromisoformat(TODAY) - datetime.timedelta(days=35)).isoformat()
molest = [ps for m in (act["mesos"] if act else []) for ps in m["pains"] if ps.split(":")[0] >= cut]
fut = [c for c in D["allComps"] if c["date"] > TODAY]; prox = min(fut, key=lambda c: c["date"]) if fut else None
def dias(d): return (datetime.date.fromisoformat(d) - datetime.date.fromisoformat(TODAY)).days
alla = sorted([(a["date"], a["notes"]) for a in D.get("analisis", [])]); last_anal = alla[-1] if alla else None
BIG = Font(bold=True, size=22, color="1F3A5F"); KLB = Font(bold=True, size=9, color="5B6472")
mcell(wsP, 1, 1, LAST, f'PANEL — {D["athlete"]}', None, Font(bold=True, size=18, color=INK), L)
estado = f'Fase {cm.get("phaseInf","—")} · Semana {cw["n"]}/{n} · {act["title"] if act else "sin plan activo"} · ' + (f'Próx: {prox["name"]} {fdm(prox["date"])} ({dias(prox["date"])}d)' if prox else "sin competencia agendada en ET")
mcell(wsP, 2, 1, LAST, estado, PatternFill("solid", fgColor="EEF2F7"), Font(bold=True, size=11, color="1F3A5F"), L)
wsP.row_dimensions[1].height = 26
# FUNDAMENTOS
nb = nivel_baremo()
_anios = f'{D["anios"]["months"]//12}a {D["anios"]["months"]%12}m'
quien = f'{PROF.get("age","")} años · {PROF.get("gender","")} · {PROF.get("height","")} / {PROF.get("weight","")} · {_anios} entrenando'.replace("  ", " ")
_adh_txt = ("s/d (sin feedback)" if no_fb else (f"{adh}%" if measurable else "s/d"))
if nb: donde = f'Nivel ≈ {nb[0]["clasificacion"]} (AIYM {nb[0]["aiym"]}) · vVO₂máx {D.get("vVO2maxZone","")} (~{nb[1]:.1f} km/h) · Adherencia {_adh_txt}'
else: donde = f'Adherencia {_adh_txt} · (sin zonas para estimar nivel)'
adonde = 'Objetivos corto/medio/largo: (a cargar por el entrenador)   |   Próx. competencia: ' + (f'{prox["name"]} {fdm(prox["date"])}' if prox else 'sin agendar en ET')
r = 3
mcell(wsP, r, 1, LAST, "📌 FUNDAMENTOS DEL PLAN  —  quién · dónde está · a dónde va", PatternFill("solid", fgColor="34495E"), Font(bold=True, size=11, color="FFFFFF"), L); r += 1
_B = [(1, 3), (4, 6), (7, 9)]
_labs = ["👤 QUIÉN (perfil ET)", "📍 DÓNDE ESTÁ (diagnóstico)", "🎯 A DÓNDE VA (lo cargás vos)"]
_txts = [quien, donde, adonde]; _lf = ["D6E0EA", "D6E0EA", "FBEFD3"]; _cf = ["F4F8FB", "F4F8FB", "FFFBF0"]
for i in range(3): mcell(wsP, r, _B[i][0], _B[i][1], _labs[i], PatternFill("solid", fgColor=_lf[i]), Font(bold=True, size=9, color="1F3A5F"), C)
for i in range(3): mcell(wsP, r + 1, _B[i][0], _B[i][1], _txts[i], PatternFill("solid", fgColor=_cf[i]), Font(size=9), L)
wsP.row_dimensions[r + 1].height = 58
r += 3
# KPIs
_BOXES = [(2, 3), (4, 5), (6, 7), (8, 9)]
def box_cols(k): return _BOXES[k]
if no_fb:
    adh_kpi = ("s/d", "atleta no cargó feedback", "FCE5CD")
    srpe_kpi = ("s/d", "sin feedback este mes", "FCE5CD")
    mol_kpi = ("s/d", "sin feedback este mes", "FCE5CD")
elif not measurable:
    adh_kpi = ("s/d", "sin semanas cerradas", "E7ECF3")
    srpe_kpi = (f'{srpe}', "carga percibida", "E7ECF3")
    mol_kpi = (("SÍ" if molest else "NO"), (molest[0][:26] if molest else "sin molestias"), "F4CCCC" if molest else "D9EAD3")
else:
    adh_kpi = (f'{adh}%', "solo runs · ver gráfico ET", "D9EAD3" if adh >= 80 else ("FFF2CC" if adh >= 60 else "F4CCCC"))
    srpe_kpi = (f'{srpe}', "carga percibida", "E7ECF3")
    mol_kpi = (("SÍ" if molest else "NO"), (molest[0][:26] if molest else "sin molestias"), "F4CCCC" if molest else "D9EAD3")
KPIS = [("VOLUMEN SEMANAL", f'{vol_act} km', f'programado · máx hist {D["volMax"]["km"] if D.get("volMax") else "—"} km', "D6E4F0"),
        ("ADHERENCIA (mes)", adh_kpi[0], adh_kpi[1], adh_kpi[2]),
        ("sRPE (mes)", srpe_kpi[0], srpe_kpi[1], srpe_kpi[2]),
        ("MOLESTIAS", mol_kpi[0], mol_kpi[1], mol_kpi[2])]
for k, (lb, val, sub, fill) in enumerate(KPIS):
    c1, c2 = box_cols(k); F = PatternFill("solid", fgColor=fill)
    mcell(wsP, r, c1, c2, lb, F, KLB, C); mcell(wsP, r + 1, c1, c2, val, F, BIG, C); mcell(wsP, r + 2, c1, c2, sub, F, Font(size=8, color="5B6472"), C)
wsP.row_dimensions[r].height = 15; wsP.row_dimensions[r + 1].height = 32; wsP.row_dimensions[r + 2].height = 15
r += 4
# (El volumen programado vs realizado se analiza directamente en ET.)
# FOCO MENSUAL  (meses en FILAS, atributos en COLUMNAS: más datos por mes, menos largo)
mcell(wsP, r, 1, LAST, "🗓️ FOCO MENSUAL  —  el QUÉ: contexto secuencial (cada mes una fila; incluye próximos)", PatternFill("solid", fgColor="34495E"), Font(bold=True, size=11, color="FFFFFF"), L); r += 1
_curk = mk(_cy, _cm)
def comps_month(y, mo):
    k = mk(y, mo); return " / ".join(f'{fdm(c["date"])} {c["name"]}' for c in D["allComps"] if (c["date"] or "")[:7] == k)
def tests_month(y, mo):
    k = mk(y, mo); return "\n".join(fdm(t["date"]) for t in D["allTests"] if (t["date"] or "")[:7] == k)
def volreal_month(y, mo):
    k = mk(y, mo); kk = [w["km"] for p in D["plans"] for w in p["weeks"] if (w["start"] or "")[:7] == k and w["km"]]
    return f'~{round(sum(kk)/len(kk))} km/sem' if kk else ""
def phase_month(y, mo):
    k = mk(y, mo)
    for w in (act["weeks"] if act else []):
        if (w.get("start") or "")[:7] == k or (w.get("end") or "")[:7] == k:
            return mbn.get(w["meso"], {}).get("phaseInf", "") or ""
    return ""
# encabezado de columnas (atributos)
_FCOLS = [("Mes", 1, "h"), ("Competencias", 2, "a"), ("Tipo / Fase", 3, "a"), ("Comentarios del atleta", 4, "a"),
          ("Carreras (ritmo real)", 5, "a"), ("Vol. sem.", 6, "a"), ("Valoraciones", 7, "a"),
          ("Obs. del entrenador", 8, "c"), ("Objetivo del mes", 9, "c")]
for name, col, own in _FCOLS:
    c = wsP.cell(r, col, name); c.fill = (COACH if own == "c" else HEAD); c.font = (LF if own == "c" else HF)
    c.alignment = C; c.border = BORD
r += 1
for (y, mo) in MONTHS:
    hl = (mk(y, mo) == _curk)
    mfill = PatternFill("solid", fgColor="C0392B") if hl else PatternFill("solid", fgColor="223B57")
    mcell(wsP, r, 1, 1, f'{_MN[mo]} {y}' + ("\n◄ HOY" if hl else ""), mfill, Font(bold=True, size=9, color="FFFFFF"), C)
    vals = [(2, comps_month(y, mo), "a"), (3, phase_month(y, mo), "a"), (4, month_synth(y, mo), "a"),
            (5, " / ".join(RACEM.get(mk(y, mo), [])), "a"), (6, volreal_month(y, mo), "a"),
            (7, tests_month(y, mo), "a"), (8, "", "c"), (9, "", "c")]
    for col, v, own in vals:
        c = wsP.cell(r, col, v); c.border = BORD; c.font = Font(size=8)
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if own == "c": c.fill = COACH
    wsP.row_dimensions[r].height = 52
    r += 1
r += 1
# VOZ + ALERTAS
mcell(wsP, r, 1, LAST, "🗣️ ÚLTIMO ANÁLISIS DEL COACH" + (f' ({fd(last_anal[0])})' if last_anal else ""), PatternFill("solid", fgColor="FFF6E6"), LF, L); r += 1
mcell(wsP, r, 1, LAST, (last_anal[1] if last_anal else "(sin análisis cargado)"), PatternFill("solid", fgColor="FFFBF0"), Font(size=9), L); wsP.row_dimensions[r].height = 110; r += 2
mcell(wsP, r, 1, LAST, "⚠ ALERTAS", None, Font(bold=True, size=11, color="9C0006"), L); r += 1
_al = []
if no_fb:
    _al.append("Sin feedback cargado este mes: adherencia / sRPE / molestias NO se pueden calcular desde ET. El atleta pudo entrenar sin registrar — verificar con él antes de asumir.")
if molest: _al += ["Molestia reciente: " + m for m in molest]
if adh < 70 and sp and not no_fb: _al.append(f"Adherencia baja ({adh}%) en el último mes")
# candidato a incrementar volumen (referencia Roces)
_disc = next((d for d in PROG["disciplinas"] if d["disciplina"] == D.get("disciplina")), None)
if _disc and D.get("volAvg") and adh >= 70 and not molest:
    _ref = _disc["por_anio_km_sem"].get(str(min(max(D["anios"]["months"] // 12, 1), 8)), "")
    _lo = int(re.match(r"(\d+)", _ref).group(1)) if _ref else 0
    if _lo and D["volAvg"] >= 0.9 * _lo:
        _al.append(f"Candidato a incrementar volumen: sostiene ~{D['volAvg']} km/sem con buena adherencia (ref. {D.get('disciplina')} año {max(D['anios']['months']//12,1)}: {_ref}). Decide el entrenador.")
if not _al: _al = ["Sin alertas activas"]
for a in _al: mcell(wsP, r, 1, LAST, "• " + a, None, Font(size=9, color="9C0006" if a != "Sin alertas activas" else "006100"), L); r += 1
wsP.sheet_view.showGridLines = False

# =================== RECOMENDACIÓN (referencias por período) ===================
wsR = wb.create_sheet("RECOMENDACIÓN")
NC = 6  # 6 columnas anchas: más aire horizontal, sin texto apretado en vertical
for i, wd in enumerate([15, 15, 13, 15, 25, 25], 1): wsR.column_dimensions[get_column_letter(i)].width = wd
def tr_jdc(s):  # traducir zonas Juan del Campo -> modelo base
    for a, b in [("Z3+", "§5"), ("Z1+", "§23"), ("Z0", "§1"), ("Z2", "§4"), ("Z3", "§45"), ("Z4", "§6"), ("Z5", "§7"), ("Z6", "§78"), ("Z1", "§2"), ("Z7", "§8")]:
        s = re.sub(re.escape(a), b, s)
    for k, v in {"§1": "Z1(Rec)", "§2": "Z2(End)", "§23": "Z2-3", "§4": "Z4(Umb)", "§45": "Z4-5", "§5": "Z5(Supra)", "§6": "Z6(VO2)", "§7": "Z7(Gluc)", "§78": "Z7-8", "§8": "Z8(Neuro)"}.items():
        s = s.replace(k, v)
    return s
pk = D.get("currentPeriodKey", "preparacion_especial")
disc = D.get("disciplina", "")
vol = D.get("volAvg", 0)
rr = 1
mcell(wsR, rr, 1, NC, f'RECOMENDACIÓN DE REFERENCIA — {D["athlete"]}', None, TITF, L); rr += 1
mcell(wsR, rr, 1, NC, f'⚠ Período INFERIDO ({D.get("currentPhase","")}) y disciplina DETECTADA de competencias ({disc}). Confirmá vos. Todo en tu modelo base de zonas. Son referencias, decidís vos.', None, SUBF, L); wsR.row_dimensions[rr].height = 26; rr += 2
mcell(wsR, rr, 1, NC, "🎯 ADAPTACIÓN OBJETIVO DEL PERÍODO", HEAD, HF, L); rr += 1
mcell(wsR, rr, 1, NC, ADV.get(pk, ""), None, Font(size=10), L); wsR.row_dimensions[rr].height = 32; rr += 2
mcell(wsR, rr, 1, NC, "📏 VOLUMEN DE REFERENCIA (Roces) vs ACTUAL", HEAD, HF, L); rr += 1
if _disc:
    _yr = max(D["anios"]["months"] // 12, 1)
    _ref = _disc["por_anio_km_sem"].get(str(min(_yr, 8)), "")
    mcell(wsR, rr, 1, NC, f'{disc} · año {_yr}: {_ref} km/sem   ·   actual: {vol} km/sem', None, Font(size=10), L); rr += 2
mcell(wsR, rr, 1, NC, f'📊 DISTRIBUCIÓN SUGERIDA (aplicada a {vol} km/sem)', HEAD, HF, L); rr += 1
# encabezado: Método(1-2) · %(3) · km(4) · Adaptación(5-6)
for c1, c2, h in [(1, 2, "Método"), (3, 3, "%"), (4, 4, "km sugeridos"), (5, 6, "Adaptación")]:
    mcell(wsR, rr, c1, c2, h, LABF, LF, C)
rr += 1
_mobj = {m["metodo"].lower().split("(")[0].strip(): m["objetivo"] for m in MET["metodos"]}
for mth, info in DIST["periodos"].get(pk, {}).get("metodos", {}).items():
    pct = info.get("pct")
    nm = mth.replace("_", " ")
    km = (f'{round(vol*pct[0]/100,1)}-{round(vol*pct[1]/100,1)} km' if pct else info.get("nota", ""))
    pcts = (f'{pct[0]}-{pct[1]}%' if pct else "—")
    obj = next((v for k, v in _mobj.items() if k.split()[0] in nm.lower()), "")
    mcell(wsR, rr, 1, 2, nm, None, CF, L); mcell(wsR, rr, 3, 3, pcts, None, CF, C)
    mcell(wsR, rr, 4, 4, km, None, CF, C); mcell(wsR, rr, 5, 6, obj, None, CF, L)
    rr += 1
rr += 1
# sesiones de referencia (Juan del Campo) — solo media/maratón (módulo maratón)
_JPMAP = {"preparacion_general": "introduccion", "preparacion_especial": "especifico_II", "competencia": "afinamiento"}
_JPNICE = {"introduccion": "introducción / base", "especifico_I": "específico I", "especifico_II": "específico II", "afinamiento": "afinamiento"}
if disc == "21.1-42.2km":
    jp = _JPMAP.get(pk)
    jd = JDC["periodos"].get(jp) if jp else None
    jsess = ((jd.get("sesiones_ejemplo_amateur_3h") or jd.get("sesiones")) if jd else None) or []
    if jsess:
        mcell(wsR, rr, 1, NC, f'🧰 SESIONES DE REFERENCIA (Juan del Campo · {_JPNICE.get(jp, jp)}) — zonas traducidas a tu modelo base', HEAD, HF, L); rr += 1
        for s in jsess[:8]:
            mcell(wsR, rr, 1, NC, "· " + tr_jdc(s), None, CF, L); rr += 1
        rr += 1

# =================== biblioteca ET (sugerencias por nombre exacto) ===================
def _nz(s): return re.sub(r"[̀-ͯ]", "", str(s or "").lower()).strip()
_DKW = {
    "21.1-42.2km": ["marat", "media", "21k", "42k", "21.1", "42.2", "champa", "16k", "consolidando"],
    "5000-10000m": ["10k", "5k", "5000", "10000", "recreacional"],
    "800-3000m":   ["5k", "velocidad", "1500", "3000", "800", "neuro", "interval"],
}
_PKW = {
    "preparacion_general":  ["base", "fondo", "endurance", "largo", "rodaje", "zona 2", "aerobic", "aeróbic", "consolidando"],
    "preparacion_especial": ["especif", "espec", "umbral", "interval", "tempo", "noruega", "vo2", "competitiv", "modelaci", "progresion", "hiit", "zona 4", "zona 5", "zona 6", "supraumbral", "ritmo"],
    "competencia":          ["taper", "afinam", "competit", "compet", "mantenimiento"],
}
_PNICE = {"preparacion_general": "preparación general", "preparacion_especial": "preparación específica", "competencia": "competencia"}
# excluir SIEMPRE lo que no es carrera de ruta; y lo de iniciación salvo principiante real
_EXCL_ALWAYS = ["trail", "ciclismo", "cycling", "bicicleta", "mtb", "nataci", "swim", "aguas abiertas", "eliptica", "elíptica"]
_EXCL_BEGIN = ["primeras zancadas", "caminar - correr", "caminar-correr", "trote y caminata", "trote caminata", "trote - caminata", "inicio de este camino", "el comienzo", "sedentar", "iniciante", "desde cero", "desde 0"]
_is_begin = (vol and vol < 18) and (D["anios"]["months"] < 14)
def _volband(t):
    t = _nz(t).replace(" ", "")
    for rgx in [r'volsem(\d{2,3})[/\-a](\d{2,3})', r'corredores(\d{2,3})k?-(\d{2,3})', r'(\d{2,3})[/\-a](\d{2,3})k/?(?:sem|porsemana)', r'\((\d{2,3})-(\d{2,3})k', r'(\d{2,3})a(\d{2,3})k']:
        m = re.search(rgx, t)
        if m:
            a, b = int(m.group(1)), int(m.group(2)); return (min(a, b), max(a, b))
    m = re.search(r'<(\d{2,3})k?', t)  # "<48k" / "corredores<48k" => hasta 48
    if m: return (0, int(m.group(1)))
    return None
def _libmatch(tp, limit):
    dkw = _DKW.get(disc, []); pkw = _PKW.get(pk, []); v = vol
    out = []; seen = set()
    for it in LIBE["items"]:
        if it.get("type") != tp: continue
        name = it.get("name", ""); fold = it.get("folder", "")
        txt = _nz(name + " / " + fold)
        if any(x in txt for x in _EXCL_ALWAYS): continue
        if not _is_begin and any(x in txt for x in _EXCL_BEGIN): continue
        dh = [k for k in dkw if k in txt]; ph = [k for k in pkw if k in txt]
        band = _volband(name + " " + fold); bfit = None
        if band and v:
            lo, hi = band
            if lo <= v <= hi: bfit = ("in", f'~{v} dentro de {lo}-{hi}k')
            elif hi >= v * 0.65 and lo <= v * 1.5: bfit = ("near", f'~{v} vs {lo}-{hi}k')
            else: continue  # banda de volumen incompatible -> descartar
        if not dh and not ph: continue  # toda sugerencia matchea disciplina o período (volumen sólo filtra/pondera)
        sc = len(dh) * 3 + len(ph) + (4 if bfit and bfit[0] == "in" else (2 if bfit else 0))
        key = _nz(name)
        if key in seen: continue
        seen.add(key)
        out.append((sc, it, dh, ph, bfit))
    out.sort(key=lambda x: -x[0])
    return out[:limit]
mcell(wsR, rr, 1, NC, "📚 DE TU LIBRERÍA ET (buscá por el nombre EXACTO en tu biblioteca)", HEAD, HF, L); rr += 1
mcell(wsR, rr, 1, NC, f'Filtradas por disciplina ({disc or "s/d"}), período ({_PNICE.get(pk, pk)}) y volumen (~{vol} km/sem). Excluye trail/ciclismo/natación y, por tu nivel, contenido de iniciación. Candidatas para inspirarte; vos elegís y adaptás.', None, SUBF, L); wsR.row_dimensions[rr].height = 26; rr += 1
# encabezado: Tipo(1) · Nombre(2-4) · Carpeta(5) · Por qué(6)
for c1, c2, h in [(1, 1, "Tipo"), (2, 4, "Nombre a buscar en tu librería"), (5, 5, "Carpeta"), (6, 6, "Por qué")]:
    mcell(wsR, rr, c1, c2, h, LABF, LF, C)
rr += 1
_any = False
for tp, lim in [("PLAN", 3), ("MESO", 5), ("MICRO", 3), ("SESION", 5)]:
    for sc, it, dh, ph, bfit in _libmatch(tp, lim):
        _any = True
        why = []
        if dh: why.append("disciplina")
        if ph: why.append("período")
        if bfit: why.append(f"volumen ({bfit[1]})")
        reason = "coincide con " + " + ".join(why) if why else "coincide por nombre"
        mcell(wsR, rr, 1, 1, tp, None, Font(size=8, bold=True, color="1F3A5F"), C)
        mcell(wsR, rr, 2, 4, it.get("name", ""), None, CF, L)
        mcell(wsR, rr, 5, 5, it.get("folder", "") or "—", None, Font(size=8), L)
        mcell(wsR, rr, 6, 6, reason, None, Font(size=8), L)
        wsR.row_dimensions[rr].height = 28
        rr += 1
if not _any:
    mcell(wsR, rr, 1, NC, "· Sin coincidencias claras en la librería para esta combinación; revisá tus carpetas por disciplina/período manualmente.", None, CF, L); rr += 1
rr += 1
mcell(wsR, rr, 1, NC, "Nota: las zonas y ritmos se expresan SIEMPRE en el modelo base (Z1 Recuperación … Z8 Neuromuscular). Los min/km reales salen de las zonas del atleta en ET.", None, SUBF, L)

# =================== orden + guardar ===================
order = ["PANEL", "DASHBOARD", "PLANIFICACIÓN ACTUAL", "HISTORIAL (planes previos)", "PERFILADO Y COMPETENCIAS", "RECOMENDACIÓN"]
wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)
safe = re.sub(r"[^\w ]", "", D["athlete"]).strip()[:40]
out = f'report/Master-{safe}.xlsx'
wb.save(out); print(out)
