#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Informe de RENOVACIÓN — Excel con una pestaña por atleta, en vista SEMANAL:
una tabla por microciclo, columnas = días Lunes..Domingo, filas = atributos de la
sesión, y un renglón superior combinado con el rango de fechas (Lun al Dom).
Lee report/renewal-<plan>-data.json. Texto ajustado y alineado al medio. Sin fórmulas."""
import json, sys, re, unicodedata
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

files = sys.argv[1:]
if not files:
    print("uso: make_renewal_xlsx.py report/renewal-<id>-data.json ..."); sys.exit(1)
plans = [json.load(open(f)) for f in files]
def _adh(d):
    tot=sum(m["progSessions"] for m in d["micros"]); done=sum(m["doneSessions"] for m in d["micros"])
    return (done/tot) if tot else 0
plans.sort(key=_adh)

INK="1A1F2B"; MUT="5B6472"
HEADER=PatternFill("solid",fgColor="1F3A5F"); HEADERF=Font(color="FFFFFF",bold=True,size=11)
DAYF=PatternFill("solid",fgColor="2E4A6B"); LABF=PatternFill("solid",fgColor="EEF2F7")
GAP=PatternFill("solid",fgColor="FBE3E3"); OKC=PatternFill("solid",fgColor="E7F6EE"); WARN=PatternFill("solid",fgColor="FDF3E6")
REST=PatternFill("solid",fgColor="F7F8FA"); DONEF=PatternFill("solid",fgColor="F3F7F4")
TITLEF=Font(bold=True,size=15,color=INK); SUBF=Font(italic=True,size=10,color=MUT); H2F=Font(bold=True,size=12,color="1F3A5F"); BOLD=Font(bold=True,color=INK)
thin=Side(style="thin",color="D8DEE8"); BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
CATFONT={"LIT":Font(bold=True,size=9,color="1C4E8A"),"ThT":Font(bold=True,size=9,color="9A7400"),"HIT":Font(bold=True,size=9,color="B02A2A")}
PAINF=PatternFill("solid",fgColor="FBE3E3")
CEN=Alignment(horizontal="center",vertical="center",wrap_text=True)
LEFT=CEN   # requisito: todo el archivo centrado, alineado al medio y texto ajustado
DAYS=["Lunes","Martes","Miércoles","Jueves","Viernes","Sábado","Domingo"]

def adh_fill(p): return OKC if p>=70 else (WARN if p>=50 else GAP)
def W(ws,ws_w):
    for i,w in enumerate(ws_w,1): ws.column_dimensions[get_column_letter(i)].width=w
def safe_title(name, used):
    t="".join(ch for ch in name if ch not in '[]:*?/\\')[:28] or "Atleta"
    base=t; i=2
    while t in used: t=f"{base[:25]}~{i}"; i+=1
    used.add(t); return t
def pace_sec(p):
    m=re.search(r"(\d+)'(\d+)", p or ""); return int(m.group(1))*60+int(m.group(2)) if m else 9999
def work_paces(laps):
    if not laps: return ""
    core = laps[1:-1] if len(laps) > 2 else laps
    segs=[(L.get("pace") or "").replace(" ","") for L in core
          if "h" not in (L.get("pace") or "") and L.get("distKm",0)>=0.4 and pace_sec(L.get("pace"))<=600]
    uniq=sorted(set(segs), key=pace_sec)
    return (uniq[0] + (f"–{uniq[-1]}" if len(uniq)>1 else "")) if uniq else ""
def d2(s):
    try: return date.fromisoformat(s)
    except Exception: return None
def fmt(d): return d.strftime("%d/%m/%Y")

# --- Respuestas del formulario de renovación (Google Form) por atleta ---
# report/form-responses.json: { <norm>: {name, date(ISO), email, qa:[[label,answer],...]} }
# Se inyectan en la pestaña de cada atleta para no tener que abrir el form aparte.
def _ntok(s):
    s=unicodedata.normalize("NFD",s or "").encode("ascii","ignore").decode().lower()
    s=re.sub(r"\(.*?\)"," ",s)                      # saca aclaraciones entre paréntesis
    return [t for t in re.split(r"[^a-z0-9]+",s) if t]
try:
    FORMS=list(json.load(open("report/form-responses.json")).values())
except Exception:
    FORMS=[]
def match_form(name):
    nt=_ntok(name); best=None; bestsc=0
    for v in FORMS:
        ft=_ntok(v.get("name")); sc=0
        for a in ft:
            if any(a==b or (len(a)>=3 and len(b)>=3 and (a.startswith(b) or b.startswith(a))) for b in nt):
                sc+=1
        if sc>bestsc: bestsc=sc; best=v
    return best if bestsc>=2 else None
def _pain_answer(ans):
    a=unicodedata.normalize("NFD",ans or "").encode("ascii","ignore").decode().lower().strip()
    return bool(a) and not re.match(r"^(no|nada|ningun|sin )", a)

# filas (atributos) de cada tabla semanal
def val_sesion(s): return (("★ " if s.get("key") else "")+(s.get("cat") or "")+" "+(s.get("title") or "")).strip()
def val_rpe(s):
    a = s.get("progRpe"); b = s.get("sRPE")
    return (f'{a if a is not None else "–"} / {b if b is not None else "–"}')
def val_dur(s):
    return f'{s.get("progDur") or "–"}' + (f' → {s["realDur"]}' if s.get("realDur") else '')
def val_dist(s): return s.get("realDist") or s.get("progDist") or "–"
def val_obj(s):
    tp=[]; [tp.append(t["target"]) for t in (s.get("targets") or []) if t.get("target") and t["target"] not in tp]
    return " · ".join(tp[:4]) if tp else (s.get("targetPace") or "–")
def val_real(s):
    if s.get("source")=="FIT": return work_paces(s.get("laps") or []) or "ver laps"
    return s.get("realPace") or "–"
def val_laps(s):
    laps=s.get("laps") or []
    if laps:
        return "\n".join(f'{L["n"]:>2} {L["time"]:>6} {L["distKm"]:.2f}km {(L.get("pace") or "").replace(" ","")} {L.get("hrA","")}/{L.get("hrM","")}' for L in laps)
    if s.get("future"): return "(futura)"
    if s.get("source") in ("manual","comentario","manual+coment."): return "s/dispositivo"
    return "–"
def val_coment(s): return (s.get("comment") or "")

# --- estado "Realizada": verde=hecha como programada · amarillo=hecha pero distinto · rojo=no hecha ---
DEVIA = re.compile(r"en\s+(lugar|vez)\s+de|cambi[eé]|modifiqu|no\s+pude|no\s+lo\s+hice|no\s+complet|"
                   r"cort[eé]|fren[eé]|camin[eé]|suspend|abandon|enferm|descompuest|me\s+sent[ií]\s+mal|"
                   r"rodaje\s+suave|m[aá]s\s+suave|reduje|menos\s+(series|vueltas|repet)|la\s+mitad|"
                   r"baj[eé]\s+(la|el)|no\s+como\s+estaba|distinto\s+a\s+lo", re.I)
def is_deviation(s):
    if not s.get("executed"): return False
    if s.get("pain"): return True
    if DEVIA.search(s.get("comment") or ""): return True
    a=s.get("progRpe"); b=s.get("sRPE")
    if a is not None and b is not None and (a-b) >= 3: return True
    return False
def day_status(sessions):
    sessions=[s for s in sessions if not s.get("strength")]   # solo runs definen el estado
    due=[s for s in sessions if s.get("due",True)]
    if not due:
        return "future" if sessions else None          # futura (gris) o descanso (sin run programado)
    executed=[s for s in due if s.get("executed")]
    if not executed: return "rojo"                      # programada y no realizada
    if any(is_deviation(s) for s in executed) or len(executed) < len(due): return "amarillo"
    return "verde"
STATUS_FILL={"verde":PatternFill("solid",fgColor="C6EFCE"),"amarillo":PatternFill("solid",fgColor="FFEB9C"),
             "rojo":PatternFill("solid",fgColor="FFC7CE"),"future":PatternFill("solid",fgColor="EDEFF2")}
STATUS_FONT={"verde":Font(bold=True,size=9,color="006100"),"amarillo":Font(bold=True,size=9,color="9C6500"),
             "rojo":Font(bold=True,size=9,color="9C0006"),"future":Font(size=9,color="5B6472")}
STATUS_TXT={"verde":"Realizada ✓","amarillo":"Modificada ~","rojo":"No realizada ✗","future":"Futura"}

ROWS=[("Sesión",val_sesion),("RPE p · sRPE",val_rpe),("Duración p→r",val_dur),("Distancia",val_dist),
      ("Ritmo objetivo",val_obj),("Ritmo real",val_real),("Laps / detalle",val_laps),("Comentario",val_coment)]

def cell_join(sessions, fn):
    parts=[]
    for s in sessions:
        v=fn(s)
        if v in (None,"","–"): continue
        parts.append((f'{s.get("moment","")}: ' if len(sessions)>1 else '')+str(v))
    return "\n".join(parts)

wb=Workbook(); used=set()

# ---- Resumen ----
ws=wb.active; ws.title="Resumen"
ws["A1"]="Informe de renovación"; ws["A1"].font=TITLEF; ws["A1"].alignment=LEFT
ws["A2"]=f"{len(plans)} atleta(s) · último mesociclo del plan vigente · adherencia sobre sesiones vencidas · generado {plans[0].get('generatedAt','')}"; ws["A2"].font=SUBF; ws["A2"].alignment=LEFT
W(ws,[26,30,16,13,12,36]); r=4
for i,h in enumerate(["Atleta","Plan","Período","Adherencia","Sesiones","Nota"],1):
    c=ws.cell(r,i,h); c.fill=HEADER; c.font=HEADERF; c.border=BORD; c.alignment=CEN
r+=1
for d in plans:
    tot=sum(m["progSessions"] for m in d["micros"]); done=sum(m["doneSessions"] for m in d["micros"])
    adh=round(100*done/tot) if tot else 0
    nota = "Sin ejecución en el período (lesión/inactividad/sin reporte)" if done==0 else ("Adherencia baja (<70%)" if adh<70 else "Adherencia buena")
    ws.cell(r,1,d["athlete"]).alignment=LEFT; ws.cell(r,2,d["plan"]["title"]).alignment=LEFT
    ws.cell(r,3," + ".join(str(x) for x in d["periodMesos"])).alignment=CEN
    c=ws.cell(r,4,f"{adh}%"); c.fill=adh_fill(adh); c.alignment=CEN
    ws.cell(r,5,f"{done}/{tot}").alignment=CEN; ws.cell(r,6,nota).alignment=LEFT
    for cc in range(1,7): ws.cell(r,cc).border=BORD
    r+=1

# ---- una pestaña por atleta: tablas semanales ----
NCOL=8   # 1 etiqueta + 7 días
for d in plans:
    ws=wb.create_sheet(safe_title(d["athlete"], used))
    p=d["plan"]
    # nombre de atleta y nombre del plan: celdas combinadas al ancho de la tabla (A..H)
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=NCOL)
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=NCOL)
    ws["A1"]=d["athlete"]; ws["A1"].font=TITLEF; ws["A1"].alignment=CEN
    ws["A2"]=f'{p["title"]} · {p["start"]} → {p["end"]} · {p["status"]} · período: {" + ".join(str(x) for x in d["periodMesos"])}   —   (la adherencia considera solo los runs; no se reportan feedbacks de fuerza)'; ws["A2"].font=SUBF; ws["A2"].alignment=CEN
    W(ws,[15]+[30.86]*7); r=4
    # --- bloque: respuestas del formulario de renovación (si las hay) ---
    fr=match_form(d["athlete"])
    if fr:
        fd=d2(fr.get("date")); fdtxt=fmt(fd) if fd else (fr.get("date") or "")
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=NCOL)
        hc=ws.cell(r,1,f'Formulario de renovación — respondido {fdtxt}')
        hc.fill=HEADER; hc.font=HEADERF; hc.alignment=CEN
        for c in range(1,NCOL+1): ws.cell(r,c).border=BORD
        r+=1
        for lab,ans in fr.get("qa",[]):
            lc=ws.cell(r,1,lab); lc.fill=LABF; lc.font=Font(bold=True,size=9,color="1F3A5F"); lc.alignment=CEN; lc.border=BORD
            ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=NCOL)
            ac=ws.cell(r,2,ans); ac.font=Font(size=9); ac.border=BORD
            ac.alignment=Alignment(horizontal="left",vertical="center",wrap_text=True)
            if lab=="Dolores/molestias" and _pain_answer(ans):
                ac.fill=PAINF; ac.font=Font(size=9,bold=True,color="B02A2A")
            for c in range(3,NCOL+1): ws.cell(r,c).border=BORD
            r+=1
        r+=1   # respiro antes de las tablas semanales
    for m in d["micros"]:
        # lunes..domingo de la semana del micro
        base = d2(m.get("start"))
        if base is None:
            ds=[d2(s["date"]) for s in m["sessions"] if d2(s["date"])]
            base = min(ds) if ds else None
        if base is None: continue
        monday = base - timedelta(days=base.weekday())
        week = [monday + timedelta(days=i) for i in range(7)]
        byday = {w: [] for w in week}
        for s in m["sessions"]:
            sd=d2(s["date"])
            if sd in byday: byday[sd].append(s)
            elif sd is not None:  # fuera de la semana (raro): a la columna más cercana
                byday[min(week, key=lambda w: abs((w-sd).days))].append(s)
        # fila título combinada
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=NCOL)
        tc=ws.cell(r,1,f'Microciclo {m["number"]}   ·   {fmt(monday)} al {fmt(week[-1])}   ·   {m["doneSessions"]} sesiones ejecutadas vs {m["progSessions"]} programadas   ({m["adherence"]}%)')
        tc.fill=HEADER; tc.font=HEADERF; tc.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        for c in range(1,NCOL+1): ws.cell(r,c).border=BORD
        r+=1
        # encabezado de días
        hc=ws.cell(r,1,""); hc.fill=LABF; hc.border=BORD
        for i,w in enumerate(week):
            c=ws.cell(r,2+i,f'{DAYS[i]}\n{w.strftime("%d/%m")}'); c.fill=DAYF; c.font=HEADERF; c.alignment=CEN; c.border=BORD
        r+=1
        # fila "Realizada" (semáforo por día)
        lc=ws.cell(r,1,"Realizada"); lc.fill=LABF; lc.font=Font(bold=True,size=9,color="1F3A5F"); lc.alignment=CEN; lc.border=BORD
        for i,w in enumerate(week):
            st=day_status(byday[w])
            c=ws.cell(r,2+i,STATUS_TXT.get(st,"") if st else "")
            if st: c.fill=STATUS_FILL[st]; c.font=STATUS_FONT[st]
            else: c.fill=REST
            c.alignment=CEN; c.border=BORD
        r+=1
        # filas de atributos
        for label, fn in ROWS:
            lc=ws.cell(r,1,label); lc.fill=LABF; lc.font=Font(bold=True,size=9,color="1F3A5F"); lc.alignment=CEN; lc.border=BORD
            for i,w in enumerate(week):
                sess=byday[w]
                txt = cell_join(sess, fn) if sess else ""
                c=ws.cell(r,2+i,txt)
                c.alignment = CEN
                c.font=Font(size=9)
                if label=="Laps / detalle": c.font=Font(size=8)
                if not sess: c.fill=REST
                if label=="Sesión":
                    cat = sess[0].get("cat") if sess else None
                    if cat in CATFONT: c.font=CATFONT[cat]
                c.border=BORD
            r+=1
        r+=1   # espacio entre micros

out="report/Renovacion.xlsx"; wb.save(out); print("OK ->",out)
